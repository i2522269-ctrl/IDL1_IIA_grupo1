# -*- coding: utf-8 -*-
# Generador del informe v5: incluye análisis de escalabilidad (14.6, 15.4)
# y recomendaciones ampliadas (R8–R12). Basado en estructura del v4.
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

RUTA_OUT = "Informe_Proyecto_Productivo_IIA_v5.docx"

def add_title(doc, text, level=1):
    p = doc.add_paragraph(text)
    try:
        if level == 1:
            p.style = 'Heading 1'
        elif level == 2:
            p.style = 'Heading 2'
        elif level == 3:
            p.style = 'Heading 3'
    except Exception:
        pass
    return p

def add_para(doc, text):
    doc.add_paragraph(text)

def add_bullet(doc, text):
    try:
        doc.add_paragraph(text, style='List Bullet')
    except Exception:
        doc.add_paragraph(f"- {text}")

def add_table(doc, headers, rows):
    t = doc.add_table(rows=len(rows)+1, cols=len(headers))
    try:
        t.style = 'Table Grid'
    except Exception:
        pass
    hdr = t.rows[0].cells
    for i,h in enumerate(headers):
        hdr[i].text = str(h)
    for r,vals in enumerate(rows):
        cells = t.rows[r+1].cells
        for i,v in enumerate(vals):
            cells[i].text = str(v)

def main():
    doc = Document()
    # Portada
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Proyecto Productivo IIA")
    run.bold = True
    run.font.size = Pt(20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Análisis predictivo de inventario y compras")
    run.font.size = Pt(14)
    doc.add_page_break()

    # 13
    add_title(doc, "13. Establecimiento de métricas", 1)
    add_para(doc, "Las métricas del proyecto se definen en un diccionario único y trazable (METRICAS.md), con nombre técnico, definición, fórmula, fuente y uso. Se sigue una convención de prefijos (total_, lead_time_, frecuencia_, precio_, R2/MAE/RMSE/MAPE) para mantener un lenguaje común entre ETL, dashboard e informe.")
    add_para(doc, "Métricas operativas: gold.metricas_prediccion_proveedor (atenciones, lead time promedio, desviación de lead time, volumen total suministrado) y gold.prediccion_demanda_mensual (cantidad despachada, frecuencia de pedidos, precio promedio pactado, último stock registrado).")
    add_para(doc, "Métricas de precisión: R², MAE, RMSE y MAPE. Criterio de validación: un modelo cuyo MAE no supera al naive no se presenta como logro. Las métricas del dashboard se calculan con base en normalizado/gold y nunca se inventan en el aire. Regla de gobernanza: si cambia el nombre o fórmula de una métrica, se actualiza METRICAS.md en el mismo commit.")
    doc.add_page_break()

    # 14
    add_title(doc, "14. Comparación de línea de base vs analítica", 1)

    add_title(doc, "14.1. Comparación técnica del enfoque", 2)
    headers = ["Dimensión", "Línea de base", "Enfoque analítico (implementado)"]
    rows = [
        ["Arquitectura de datos", "Reportes estáticos (Excel), cálculos dispersos", "Arquitectura Medallion (bronze → silver → normalizado → gold), trazable por capas"],
        ["Calidad y trazabilidad", "Origen de datos poco auditable", "ETL por fases con controles de calidad, conteos y evidencias 01–11"],
        ["Integridad de datos", "Duplicados, nulos, redundancias", "Normalización 3FN con dimensiones y hechos + FKs, análisis de redundancia e integridad referencial"],
        ["Reproducibilidad", "Resultados no deterministas", "Determinista (semilla 42), validación con conteos exactos (0 diferencias) entre gold y consultas"],
        ["Validación de modelos", "Visual/subjetiva", "Backtest riguroso frente a modelo naive (MAE/RMSE/MAPE). Solo se presenta modelo que supere al naive"],
        ["Calibración", "No sistemática", "Calibración explícita del Random Forest (grid + validación por artículo y agregada)"],
        ["Casos atípicos (huérfanos)", "Ignorados o forzados", "Decisión justificada: demanda desde silver.recepcion para conservar huerfanos (957 sin pedido + 1 fuera de catálogo) sin romper COUNT(DISTINCT solnum)"],
        ["Evidencias", "Ausentes", "Cadena de evidencia completa 01–11, verificable y reproducible"],
    ]
    add_table(doc, headers, rows)

    add_title(doc, "14.2. Comparación económica y de eficiencia", 2)
    headers = ["Indicador", "Línea de base", "Propuesta analítica"]
    rows = [
        ["Tiempo de reproceso", "Manual, horas, propenso a errores", "Reproducible/automatizable (scripts ordenados). Ingesta incremental evita recarga completa"],
        ["Costos de calidad", "Errores detectados tarde", "Detección temprana por fases + validaciones cruzadas con 0 diferencias"],
        ["Eficiencia operativa", "Basada en histórico estático o intuición", "Pronósticos a 6 meses + RF de precios (por artículo y tendencia) para anticipar compras"],
        ["Escalabilidad/transferibilidad", "Difícil reutilizar", "Reutilizable con ajustes menores (mapeo CSV + credenciales). Guía REUSO.md paso a paso"],
        ["Mantenimiento", "Conocimiento tácito", "Conocimiento explícito (ETL+SQL+evidencias+diccionario + Handoffs)"],
    ]
    add_table(doc, headers, rows)

    add_title(doc, "14.3. Impacto en la capacidad institucional y legitimidad social", 2)
    headers = ["Dimensión", "Línea de base", "Propuesta analítica"]
    rows = [
        ["Capacidad institucional", "Fragmentada, difícil transferir", "Transferencia estructurada (Handoffs + REUSO.md)"],
        ["Gobierno y lenguaje común", "Interpretaciones distintas", "Diccionario único de métricas (fórmula + fuente)"],
        ["Legitimidad técnica/social", "Poca auditabilidad → baja confianza", "Cadena de verificación 01–11 con hechos (0 diferencias, backtest, calibración)"],
        ["Transparencia metodológica", "Implícita", "Explícita (SQL versionado + ETL + decisiones técnicas justificadas)"],
        ["Reutilización y capacitación", "Dependencia de experto", "Scripts reproducibles + guía de despliegue → menor riesgo, mayor autonomía"],
        ["Ética y seguridad", "Riesgo de credenciales en código", "Credenciales fuera del código (.env, .streamlit/secrets.toml), ignorados por Git"],
    ]
    add_table(doc, headers, rows)

    add_title(doc, "14.4. Síntesis de comparación", 2)
    add_para(doc, "El salto es cualitativo: de descriptivo (qué pasó) a analítico-predictivo (qué puede pasar y con qué grado de error); de cálculos no trazables a pipeline reproducible con cadena de evidencia completa; de intuición a decisiones con métricas validadas frente a naive; de conocimiento tácito a capacidad institucional explícita y transferible. El proyecto no solo implementa modelos, sino que institucionaliza calidad de datos y reproducibilidad.")

    add_title(doc, "14.5. Argumento final de comparación", 2)
    add_para(doc, "La propuesta se justifica porque convierte datos operativos en información confiable, auditable y reutilizable. La evidencia demuestra igualdad exacta (0 diferencias) entre lo materializado en Gold y lo esperado (08_reproduccion_gold.txt, 09_resultados_gold.txt), valida pronósticos frente a línea de base predictiva (10_validacion_pronosticos.txt) y calibra modelos de precios sistemáticamente (11_calibracion_rf.txt). El diferencial es legitimidad técnica, trazabilidad y capacidad de auditoría.")

    add_title(doc, "14.6. Escalabilidad del proyecto", 2)
    add_para(doc, "El proyecto es escalable por evolución (no por reescritura). Ejes clave:")
    add_bullet(doc, "Datos: arquitectura por capas + conservación de huérfanos. Escala bien hasta decenas de millones con particionamiento de hechos (preparado).")
    add_bullet(doc, "Arquitectura: Medallion + 3FN (dimensiones/hechos) permite añadir fuentes/procesos sin romper capas existentes.")
    add_bullet(doc, "ETL: ingesta incremental (checkpoint), CTEs/window functions, validaciones 08/09. Modo full por defecto garantiza reproducibilidad; incremental (upsert) propuesto manteniendo equivalencia 0 diferencias.")
    add_bullet(doc, "Computacional (ML): RF por artículo paralelizable, separación entrenamiento/inferencia recomendada (entrenamiento offline + artefactos versionados con metadatos y semilla 42).")
    add_bullet(doc, "Operativa: orquestación por fases (run_etl_completo.ps1), idempotente y verificable. Preparado para Task Scheduler/cron.")
    add_bullet(doc, "Aplicación/despliegue: Streamlit solo lectura sobre gold/normalizado, app_informe lee docx único, pre-agregados en gold para rendimiento.")
    add_bullet(doc, "Organizacional: reutilizable (REUSO.md), transferible (Handoffs), multi-tenencia conceptual (dim_empresa) con compatibilidad hacia atrás.")
    add_bullet(doc, "Riesgos/mitigantes: evitar recreación completa en prod (upsert incremental), entrenar offline, particionar por fechas/FKs cuando supere ~5–10M filas.")
    add_para(doc, "Roadmap incremental (Fase 0–5): línea base → orquestación → gold incremental (equivalencia obligatoria) → ML offline + artefactos → particionamiento/índices (cuando volumen lo requiera) → multi-tenencia. Conclusión: escalable por evolución, con reproducibilidad y trazabilidad demostradas.")
    doc.add_page_break()

    # 15
    add_title(doc, "15. Análisis económico de la propuesta a implementar", 1)
    add_para(doc, "Análisis costo-beneficio predominantemente operativo e institucional (coherente con Proyecto Productivo IIA).")

    add_title(doc, "15.1. Costos estimados", 2)
    headers = ["Concepto", "Estimación", "Observación"]
    rows = [
        ["Infraestructura", "Bajo", "PostgreSQL/Supabase existente. Incidencia fue de disponibilidad, no de diseño/costo"],
        ["Desarrollo e implementación", "Puntual (ya realizado)", "Pipeline ETL + validaciones + evidencias + app funcionales"],
        ["Operación y mantenimiento", "Muy bajo", "Scripts + ingesta incremental. Automatizable con cron/Task Scheduler"],
        ["Formación/capacitación", "Bajo", "Documentación (README, METRICAS, REUSO, Handoffs) reduce curva de aprendizaje"],
    ]
    add_table(doc, headers, rows)

    add_title(doc, "15.2. Beneficios estimados", 2)
    headers = ["Beneficio", "Tipo", "Explicación"]
    rows = [
        ["Reducción de retrabajo", "Económico", "Validaciones por fases + 0 diferencias evitan rehacer cálculos/informes"],
        ["Mejora en planeación de compras", "Económico/operativo", "Pronósticos + RF de precios permiten anticipar, mitigar sobrestock/rupturas"],
        ["Menor riesgo de error", "Económico/estratégico", "Validación vs naive + calibración sistemática"],
        ["Preservación de capital intelectual", "Institucional", "Conocimiento explícito (no tácito). Reduce dependencia de experto"],
        ["Reutilización (escala)", "Económico a futuro", "Aplicable a otra empresa/proceso con ajustes menores (amortiza inversión)"],
        ["Auditoría y cumplimiento", "Organizacional", "Cadena de evidencia 01–11 facilita auditoría y justificación"],
        ["Mayor adopción (legitimidad)", "Social", "Diccionario único + evidencia generan confianza entre áreas"],
    ]
    add_table(doc, headers, rows)

    add_title(doc, "15.3. Balance costo-beneficio", 2)
    add_para(doc, "Relación costo-beneficio favorable. Costos mayormente puntuales (ya incurridos) y beneficios recurrentes (eficiencia, calidad, trazabilidad, capacidad institucional). El mayor retorno es institucional (transferibilidad + gobernanza de datos).")

    add_title(doc, "15.4. Escalabilidad y economía de crecimiento", 2)
    add_para(doc, "La economía de crecimiento es mayoritariamente de reutilización:")
    add_bullet(doc, "Costos recurrentes vs puntuales: operación baja (incremental+orquestado), crecimiento de volumen no exige reescritura.")
    add_bullet(doc, "Beneficios recurrentes: eficiencia operativa, trazabilidad, preservación de capital intelectual y reducción de riesgo de error.")
    add_bullet(doc, "Economías de reutilización: arquitectura reutilizable tal cual; cambian solo mapeo CSV y credenciales (REUSO.md). Portabilidad a otra empresa/proceso multiplica retorno de inversión.")
    add_bullet(doc, "ROI institucional: conocimiento explícito + handoffs reducen dependencia de experto, facilitan transferencia y mitigan riesgo de pérdida de capacidad.")
    add_bullet(doc, "Escalado controlado: mejoras incrementales (gold incremental, ML offline, particionamiento cuando corresponda) permiten crecer sin romper reproducibilidad, evitando costos de reingeniería.")
    doc.add_page_break()

    # 16
    add_title(doc, "16. Conclusiones y recomendaciones", 1)

    add_title(doc, "16.1. Conclusiones", 2)
    add_bullet(doc, "Enfoque analítico riguroso: bronze→silver→normalizado (3FN)→gold con validaciones por etapas (evidencias 01–07).")
    add_bullet(doc, "Reproducibilidad demostrada empíricamente: 0 diferencias entre gold regenerado y consultas esperadas (08_reproduccion_gold.txt, 09_resultados_gold.txt).")
    add_bullet(doc, "Modelos validados (no asumidos): backtest frente a modelo naive (10_validacion_pronosticos.txt) + calibración sistemática del RF (11_calibracion_rf.txt).")
    add_bullet(doc, "Gobierno de métricas sólido: METRICAS.md como fuente única, lenguaje común entre dashboard/ETL/análisis.")
    add_bullet(doc, "Solución transferible y sostenible: REUSO.md + Handoffs + evidencias aseguran conocimiento explícito (no tácito).")
    add_bullet(doc, "Decisiones técnicas justificadas con datos: lectura de demanda desde silver.recepcion para conservar huerfanos (transparente y metodológicamente honesto).")
    add_bullet(doc, "Entregable integral: ETL verificado, modelos, dashboard (app.py), informe navegable (app_informe.py) y cadena completa de evidencias (01–11).")
    add_bullet(doc, "Escalable por evolución: preparado para crecer con refinamientos incrementales, sin reescritura del núcleo.")
    add_bullet(doc, "Escalabilidad preserva reproducibilidad: mejoras propuestas mantienen trazabilidad, validación vs naive y calibración RF como barreras de calidad obligatorias.")

    add_title(doc, "16.2. Recomendaciones", 2)
    add_bullet(doc, "R1. Mantener validación obligatoria vs naive (backtest) al reentrenar con nueva data.")
    add_bullet(doc, "R2. Usar ingesta incremental en operación (etl_incremental.py) + automatizar con cron/Task Scheduler.")
    add_bullet(doc, "R3. Mantener calibración del RF cuando cambie el histórico (volumen significativo). Documentar resultados.")
    add_bullet(doc, "R4. Preservar cadena de evidencias 01–11: renovar evidencias correspondientes al regenerar capas (no arrastrar números antiguos).")
    add_bullet(doc, "R5. Conservar METRICAS.md como fuente de verdad: cualquier cambio de nombre/fórmula actualizarlo en el mismo commit.")
    add_bullet(doc, "R6. Aplicar enfoque a nuevos casos siguiendo checklist REUSO.md (credenciales, mapeo CSV, ETL ordenado, evidencias renovadas, modelos con semilla 42).")
    add_bullet(doc, "R7. Actualizar nota de infraestructura en informe cuando Supabase se recupere, manteniendo registro del backup y trazabilidad.")
    add_bullet(doc, "R8. Implementar orquestación por fases (run_etl_completo.ps1) para ejecución ordenada y repetible.")
    add_bullet(doc, "R9. Adoptar gold incremental (upsert) manteniendo modo full por defecto y exigiendo equivalencia 0 diferencias entre ambos modos antes de usarlo en producción.")
    add_bullet(doc, "R10. Migrar a entrenamiento ML offline + artefactos versionados (con metadatos, fecha, semilla 42, versión de dataset) con fallback a entrenamiento bajo demanda. Registrar en evidencia 12_modelos_versionados.txt.")
    add_bullet(doc, "R11. Evaluar particionamiento e índices por fechas/FKs cuando el volumen supere ~5–10M filas, sin alterar lógica de negocio.")
    add_bullet(doc, "R12. Considerar multi-tenencia por empresa_id (dim_empresa) si escala a múltiples unidades/empresas, garantizando compatibilidad hacia atrás.")

    doc.save(RUTA_OUT)
    print(f"Informe generado: {RUTA_OUT}")

if __name__ == "__main__":
    main()
